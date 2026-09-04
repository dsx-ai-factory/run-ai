import re
from dataclasses import dataclass
from collections import defaultdict
from typing import Dict
from field import Field, Variant
from target import Target
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

type_to_letter = {"TemplateSpec" : "T", "Spec": "S", "Rules": "R", "Defaults": "D"}

@dataclass
class Config:
    fields: Dict[str, Field]
    targets: Dict[str, Target]

    @classmethod
    def from_dict(cls, data: Dict, base_dir: str) -> "Config":
        if "targets" not in data or "fields" not in data:
            raise ValueError("Config must have both 'targets' and 'fields' sections.")

        targets = {}
        for target_id, t in data["targets"].items():
            targets[target_id] = Target(
                id=target_id,
                type=t.get("type"),
                path=t.get("path"),
                skeleton=t.get("skeleton"),
                comment=t.get("comment"),
                nullable=t.get("nullable"),
                name_change=t.get("nameChange", {}),
                extract_fields=t.get("extractFields", None)
            )
            targets[target_id].validate(base_dir)

        used_names = set()
        fields = {}

        for field_id, f in data["fields"].items():
            if "name" not in f:
                raise ValueError(f"Field '{field_id}' missing 'name' key.")

            field_name = f["name"]

            # Check for uniqueness of the 'name' across all fields
            if field_name in used_names:
                existing_id = next(k for k, v in fields.items() if v.name == field_name)
                raise ValueError(f"Duplicate field name '{field_name}' used in both '{existing_id}' and '{field_id}'.")

            used_names.add(field_name)

            variants = [Variant(**v) for v in f.get("variants", [])]

            field_obj = Field(
                id=field_id,
                name=f["name"],
                variants=variants)
            field_obj.validate(targets)
            fields[field_id] = field_obj

        return cls(fields=fields, targets=targets)

    def export_excel(self, path: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Fields Coverage"

        ws.freeze_panes = "C2"

        # target ending with Spec, Rules or Defaults needs to be grouped in the same column
        # all the rest will get one per column
        grouped_targets = defaultdict(dict)

        for target_id in self.targets:
            match = re.match(r"^(.*?)(TemplateSpec|Spec|Rules|Defaults)$", target_id)
            if match:
                base, typ = match.groups()
                grouped_targets[base][typ] = target_id
            else:
                grouped_targets[target_id] = {}  # targets that don't match pattern

        header = ["Field Name", "Variant Name"] + [humanize_target_name(tid) for tid in list(grouped_targets.keys())]
        ws.append(header)

        wrap_alignment = Alignment(wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center")

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = wrap_alignment

        for field in self.fields.values():
            for variant in field.variants:
                field_name = re.sub(r"[.]", " -> ", field.name)

                row = [field_name, variant.name]

                for base in grouped_targets:
                    if grouped_targets.get(base):
                        letters = ""
                        for typ, tid in grouped_targets[base].items():
                            if tid in variant.targets:
                                letters += type_to_letter[typ]
                        row.append(letters)
                    else:
                        row.append("X" if base in variant.targets else "")

                ws.append(row)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.alignment = wrap_alignment

        ws.column_dimensions["A"].width = 30

        wb.save(path)

def humanize_target_name(name: str) -> str:
    return re.sub(r"(?<!^)(?=[A-Z])", " ", name)